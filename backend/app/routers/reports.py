import csv
import os
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app import models


router = APIRouter(
    prefix="/reports",
    tags=["reports"]
)


# ============================================================
# REPORTS DIRECTORY
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.dirname(
            os.path.abspath(__file__)
        )
    )
)

REPORTS_DIR = os.path.join(
    BASE_DIR,
    "reports"
)

os.makedirs(
    REPORTS_DIR,
    exist_ok=True
)


# ============================================================
# GET LATEST ANALYSIS FOR EACH CAMERA
# ============================================================

def get_latest_analysis_per_camera(db: Session):

    cameras = (
        db.query(models.Camera)
        .order_by(models.Camera.id.asc())
        .all()
    )

    rows = []

    for camera in cameras:

        analysis = (
            db.query(models.AnalysisResult)
            .filter(
                models.AnalysisResult.camera_id == camera.id
            )
            .order_by(
                models.AnalysisResult.scan_started_at.desc().nullslast(),
                models.AnalysisResult.created_at.desc()
            )
            .first()
        )

        if analysis:

            rows.append({
                "camera": camera.name,
                "camera_id": camera.id,
                "location": camera.location or "",
                "capacity": camera.capacity or 500,
                "filename": analysis.filename or "",
                "scan_time": (
                    analysis.scan_started_at
                    or analysis.created_at
                ),
                "people_count": analysis.people_count or 0,
                "average_people": analysis.average_people or 0,
                "crowd_density": analysis.crowd_density or 0,
                "risk_level": (
                    analysis.risk_level or "LOW"
                ).upper(),
                "detection_count": len(
                    analysis.boxes or []
                ),
            })

    return rows


# ============================================================
# CREATE CSV REPORT
# ============================================================

def create_csv_report(rows):

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        f"crowd_report_{timestamp}.csv"
    )

    filepath = os.path.join(
        REPORTS_DIR,
        filename
    )

    fieldnames = [
        "Camera",
        "Camera ID",
        "Location",
        "Capacity",
        "Scan Time",
        "Video File",
        "People Count",
        "Average People",
        "Crowd Density (%)",
        "Risk Level",
        "Detection Count",
    ]

    with open(
        filepath,
        "w",
        newline="",
        encoding="utf-8"
    ) as file:

        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames
        )

        writer.writeheader()

        for row in rows:

            writer.writerow({
                "Camera": row["camera"],
                "Camera ID": row["camera_id"],
                "Location": row["location"],
                "Capacity": row["capacity"],
                "Scan Time": (
                    row["scan_time"].isoformat()
                    if row["scan_time"]
                    else ""
                ),
                "Video File": row["filename"],
                "People Count": row["people_count"],
                "Average People": row["average_people"],
                "Crowd Density (%)": row["crowd_density"],
                "Risk Level": row["risk_level"],
                "Detection Count": row["detection_count"],
            })

    return filename


# ============================================================
# CREATE POWERPOINT REPORT
# ============================================================

def create_pptx_report(rows):

    from pptx import Presentation
    from pptx.util import Inches, Pt

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        f"crowd_report_{timestamp}.pptx"
    )

    filepath = os.path.join(
        REPORTS_DIR,
        filename
    )

    presentation = Presentation()

    # --------------------------------------------------------
    # TITLE SLIDE
    # --------------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )

    slide.shapes.title.text = (
        "AI Crowd Risk Analysis Report"
    )

    slide.placeholders[1].text = (
        "Generated: "
        + datetime.now().strftime(
            "%d %B %Y, %H:%M:%S"
        )
    )

    # --------------------------------------------------------
    # SUMMARY SLIDE
    # --------------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[1]
    )

    slide.shapes.title.text = (
        "Latest YOLO Analysis"
    )

    total_people = sum(
        int(row["people_count"])
        for row in rows
    )

    high_risk = sum(
        1
        for row in rows
        if row["risk_level"] in [
            "HIGH",
            "CRITICAL"
        ]
    )

    slide.placeholders[1].text = (
        f"Analyzed cameras: {len(rows)}\n"
        f"Total detected people: {total_people}\n"
        f"High/Critical risk cameras: {high_risk}\n"
        f"Generated: "
        f"{datetime.now().strftime('%d-%m-%Y %H:%M:%S')}"
    )

    # --------------------------------------------------------
    # CAMERA DETAILS
    # --------------------------------------------------------

    for row in rows:

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[5]
        )

        slide.shapes.title.text = (
            f"{row['camera']} - Crowd Analysis"
        )

        textbox = slide.shapes.add_textbox(
            Inches(1),
            Inches(1.5),
            Inches(8),
            Inches(4.5)
        )

        text_frame = textbox.text_frame

        information = [
            f"Location: {row['location']}",
            f"Capacity: {row['capacity']}",
            f"Video: {row['filename']}",
            f"Scan Time: {row['scan_time']}",
            f"People Count: {row['people_count']}",
            f"Average People: {row['average_people']}",
            f"Crowd Density: {row['crowd_density']}%",
            f"Risk Level: {row['risk_level']}",
            f"Detection Count: {row['detection_count']}",
        ]

        text_frame.text = information[0]

        for item in information[1:]:

            paragraph = text_frame.add_paragraph()
            paragraph.text = item
            paragraph.font.size = Pt(18)

    presentation.save(filepath)

    return filename


# ============================================================
# GENERATE REPORT
# ============================================================

@router.post("/generate")
def generate_report(
    db: Session = Depends(get_db)
):

    rows = get_latest_analysis_per_camera(db)

    if not rows:

        raise HTTPException(
            status_code=404,
            detail="No YOLO analysis data available yet."
        )

    csv_filename = create_csv_report(rows)

    pptx_filename = None

    try:

        pptx_filename = create_pptx_report(
            rows
        )

    except Exception as error:

        print(
            "PPTX generation failed:",
            error
        )

    return {
        "success": True,
        "message": "Report generated successfully",
        "csv": csv_filename,
        "pptx": pptx_filename,
        "generated_at": datetime.now().isoformat(),
        "camera_count": len(rows),
    }


# ============================================================
# LIST REPORTS
# ============================================================

@router.get("")
def list_reports():

    if not os.path.exists(REPORTS_DIR):
        return []

    reports = []

    for filename in os.listdir(
        REPORTS_DIR
    ):

        filepath = os.path.join(
            REPORTS_DIR,
            filename
        )

        if not os.path.isfile(filepath):
            continue

        extension = (
            filename.rsplit(".", 1)[-1].upper()
            if "." in filename
            else "FILE"
        )

        size_bytes = os.path.getsize(
            filepath
        )

        reports.append({
            "name": filename,
            "filename": filename,
            "type": extension,
            "size": size_bytes,
            "created_at": datetime.fromtimestamp(
                os.path.getctime(filepath)
            ).isoformat(),
        })

    reports.sort(
        key=lambda item: item["created_at"],
        reverse=True
    )

    return reports


# ============================================================
# DOWNLOAD REPORT
# ============================================================

@router.get("/download/{filename}")
def download_report(
    filename: str
):

    safe_filename = os.path.basename(
        filename
    )

    filepath = os.path.join(
        REPORTS_DIR,
        safe_filename
    )

    if not os.path.exists(filepath):

        raise HTTPException(
            status_code=404,
            detail="Report not found"
        )

    return FileResponse(
        filepath,
        filename=safe_filename
    )